#!/usr/bin/env python3
"""Turns hero-balloon-styles.xlsx (the table filled in by hand) into data/hero-styles.json.
Usage: python3 tools/build-hero-styles.py path/to/hero-balloon-styles.xlsx
Locked rows (Deadpool, Dark Phoenix, Cursed form) are skipped: the app keeps their built-in look."""
import json, re, sys, warnings
import openpyxl
warnings.filterwarnings('ignore')
slug = lambda n: re.sub(r'^-|-$', '', re.sub(r'[^a-z0-9]+', '-', n.lower()))
wb = openpyxl.load_workbook(sys.argv[1])
fonts = {}
for fid, name, weight, *_ in list(wb['Fonts'].iter_rows(min_row=2, values_only=True)):
    if fid: fonts[fid] = {'family': name, 'weight': weight}
heroes, problems = {}, []
for hero, form, fid, _name, caps, line, locked, *_ in wb['Heroes'].iter_rows(min_row=2, values_only=True):
    if not hero or locked == 'כן': continue
    if fid not in fonts: problems.append(f'{hero} {form}: unknown font {fid!r}'); continue
    line = (line or '').strip()
    if line and not line.startswith('#'): line = '#' + line
    if not re.fullmatch(r'#[0-9a-fA-F]{6}', line): problems.append(f'{hero} {form}: bad color {line!r}'); continue
    st = {'font': fid, 'caps': caps == 'כן', 'line': line.lower()}
    h = heroes.setdefault(slug(hero), {})
    if form == 'Normal': h.update(st)
    else: h.setdefault('forms', {})[slug(form)] = st
if problems: sys.exit('\n'.join(problems))
json.dump({'fonts': fonts, 'heroes': heroes}, open('data/hero-styles.json', 'w'), ensure_ascii=False, indent=1)
print(len(heroes), 'heroes written')
